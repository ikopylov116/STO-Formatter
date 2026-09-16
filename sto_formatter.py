import re
from pathlib import Path
from tkinter import Tk, StringVar, filedialog, messagebox, Text, END
import tkinter as tk
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

APP_TITLE = "СТО Formatter"
DEFAULTS = {"top": 2.0, "bottom": 2.0, "left": 3.0, "right": 1.0,
            "font": "Times New Roman", "size": 14, "line": 1.5, "first": 1.25}
STRUCTURAL = {"реферат", "аннотация", "оглавление", "содержание", "введение", "заключение",
              "выводы", "нормативные ссылки", "определения, обозначения и сокращения",
              "список использованных источников", "список литературы", "приложения"}

def set_run_font(run, name, size):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), name)

def add_page_number(paragraph, font_name, font_size):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.text = ""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])
    set_run_font(run, font_name, font_size)

def numbered_heading(text):
    return bool(re.match(r"^\s*\d+(?:\.\d+)*\s+\S", text))

def is_heading(paragraph):
    text = paragraph.text.strip()
    if not text:
        return False
    plain = text.lower().rstrip(".:")
    return plain in STRUCTURAL or numbered_heading(text) or paragraph.style.name.lower().startswith("heading")

def read_template(pdf_path):
    cfg = dict(DEFAULTS); notes = []
    if not pdf_path or PdfReader is None:
        notes.append("PDF-анализ недоступен; использованы параметры профиля STO P_305.2026.")
        return cfg, notes
    try:
        reader = PdfReader(pdf_path)
        if reader.pages:
            page = reader.pages[0]
            w = float(page.mediabox.width) * 25.4 / 72
            h = float(page.mediabox.height) * 25.4 / 72
            notes.append(f"Шаблон PDF: {w:.1f} × {h:.1f} мм.")
            if abs(w - 210) < 4 and abs(h - 297) < 4:
                notes.append("Определён формат A4.")
        text = "\n".join((p.extract_text() or "") for p in reader.pages[:3]).lower()
        if "times new roman" in text:
            cfg["font"] = "Times New Roman"
        m = re.search(r"(\d{1,2})\s*(?:pt|пт|пункт)", text)
        if m and 8 <= int(m.group(1)) <= 24:
            cfg["size"] = int(m.group(1))
        if "1,5" in text or "1.5" in text:
            cfg["line"] = 1.5
        notes.append("Параметры, отсутствующие в PDF, берутся из профиля STO P_305.2026.")
    except Exception as exc:
        notes.append(f"Не удалось полностью прочитать PDF: {exc}")
    return cfg, notes

def apply_format(src, dst, cfg):
    doc = Document(src)
    for section in doc.sections:
        section.top_margin = Cm(cfg["top"]); section.bottom_margin = Cm(cfg["bottom"])
        section.left_margin = Cm(cfg["left"]); section.right_margin = Cm(cfg["right"])
        add_page_number(section.footer.paragraphs[0], cfg["font"], cfg["size"])
    changed = headings = 0
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text: continue
        heading = is_heading(paragraph)
        pf = paragraph.paragraph_format
        pf.line_spacing = cfg["line"]; pf.space_before = Pt(0); pf.space_after = Pt(0)
        pf.first_line_indent = Cm(0 if heading else cfg["first"]); pf.keep_with_next = heading
        if text.lower().rstrip(".:") in STRUCTURAL:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif heading:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in paragraph.runs:
            set_run_font(run, cfg["font"], cfg["size"])
            run.underline = False
            if heading: run.bold = True
        changed += 1; headings += int(heading)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    pf = paragraph.paragraph_format
                    pf.line_spacing = cfg["line"]; pf.space_before = Pt(0); pf.space_after = Pt(0)
                    pf.first_line_indent = Cm(0)
                    for run in paragraph.runs: set_run_font(run, cfg["font"], cfg["size"])
    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                for run in paragraph.runs: set_run_font(run, cfg["font"], cfg["size"])
    settings = doc.settings.element
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    doc.save(dst)
    return changed, headings

def check_document(path, cfg):
    doc = Document(path); results = []
    for i, s in enumerate(doc.sections, 1):
        for label, actual, expected in (("левое поле", s.left_margin.cm, cfg["left"]),
            ("правое поле", s.right_margin.cm, cfg["right"]), ("верхнее поле", s.top_margin.cm, cfg["top"]),
            ("нижнее поле", s.bottom_margin.cm, cfg["bottom"])):
            ok = abs(actual - expected) <= 0.08
            results.append((ok, f"Раздел {i}: {label} {actual:.2f} см (норма {expected:.2f})."))
    doubles = sum("  " in p.text for p in doc.paragraphs)
    wrong_font = sum(1 for p in doc.paragraphs for r in p.runs if r.font.name and r.font.name != cfg["font"])
    results.append((doubles == 0, f"Двойные пробелы: {doubles}."))
    results.append((wrong_font == 0, f"Фрагментов с другим явно заданным шрифтом: {wrong_font}."))
    return results

class App:
    def __init__(self, root):
        self.root = root; root.title(APP_TITLE); root.geometry("760x560"); root.minsize(650, 480)
        self.template = StringVar(value="PDF-шаблон не выбран"); self.work = StringVar(value="DOCX-работа не выбрана")
        self.status = StringVar(value="Выберите PDF-шаблон и DOCX-работу."); self.cfg = dict(DEFAULTS); self.build()
    def build(self):
        tk.Label(self.root, text=APP_TITLE, font=("Arial", 23, "bold")).pack(pady=(18, 2))
        tk.Label(self.root, text="Автоматическое оформление DOCX по шаблону СТО", font=("Arial", 11)).pack(pady=(0, 15))
        box = tk.Frame(self.root, bd=1, relief="groove", padx=14, pady=12); box.pack(fill="x", padx=22)
        self.file_row(box, "Шаблон СТО (PDF):", self.template, self.choose_pdf)
        self.file_row(box, "Работа (DOCX):", self.work, self.choose_docx)
        buttons = tk.Frame(self.root); buttons.pack(pady=16)
        tk.Button(buttons, text="Проверить DOCX", width=20, command=self.check).pack(side="left", padx=6)
        tk.Button(buttons, text="Оформить и сохранить", width=24, command=self.format).pack(side="left", padx=6)
        tk.Label(self.root, textvariable=self.status, anchor="w").pack(fill="x", padx=24)
        self.log = Text(self.root, height=17, wrap="word", font=("Consolas", 10)); self.log.pack(fill="both", expand=True, padx=22, pady=(8, 20))
        self.log.insert(END, "Как пользоваться:\n1. Загрузите PDF-шаблон СТО.\n2. Загрузите свою работу .docx.\n3. Нажмите «Оформить и сохранить».\n4. Выберите место для нового документа.\n\nИсходный DOCX не изменяется.\n")
    def file_row(self, parent, label, variable, command):
        row = tk.Frame(parent); row.pack(fill="x", pady=5)
        tk.Label(row, text=label, width=22, anchor="w").pack(side="left")
        tk.Label(row, textvariable=variable, anchor="w", fg="gray", wraplength=440).pack(side="left", fill="x", expand=True)
        tk.Button(row, text="Выбрать", width=12, command=command).pack(side="right")
    def choose_pdf(self):
        path = filedialog.askopenfilename(filetypes=[("PDF", "*.pdf")])
        if path:
            self.template.set(path); self.cfg, notes = read_template(path); self.log.delete("1.0", END)
            self.log.insert(END, "Шаблон загружен.\n" + "\n".join(notes) + "\n"); self.status.set("PDF-шаблон принят.")
    def choose_docx(self):
        path = filedialog.askopenfilename(filetypes=[("Word DOCX", "*.docx")])
        if path: self.work.set(path); self.status.set("DOCX-работа выбрана.")
    def check(self):
        path = self.work.get()
        if not path.lower().endswith(".docx"): messagebox.showwarning(APP_TITLE, "Сначала выберите DOCX-работу."); return
        try:
            self.log.delete("1.0", END)
            for ok, msg in check_document(path, self.cfg): self.log.insert(END, ("[OK] " if ok else "[!] ") + msg + "\n")
            self.status.set("Проверка завершена.")
        except Exception as exc: messagebox.showerror(APP_TITLE, f"Ошибка проверки:\n{exc}")
    def format(self):
        src = self.work.get()
        if not src.lower().endswith(".docx"): messagebox.showwarning(APP_TITLE, "Сначала выберите DOCX-работу."); return
        dst = filedialog.asksaveasfilename(defaultextension=".docx", initialfile=f"{Path(src).stem}_STO.docx", filetypes=[("Word DOCX", "*.docx")])
        if not dst: return
        if Path(src).resolve() == Path(dst).resolve(): messagebox.showwarning(APP_TITLE, "Выберите другое имя файла."); return
        try:
            changed, headings = apply_format(src, dst, self.cfg)
            self.log.insert(END, f"\nГотово. Обработано абзацев: {changed}; заголовков: {headings}.\nСохранено: {dst}\n")
            self.status.set("Документ по STO создан."); messagebox.showinfo(APP_TITLE, f"Готово!\n\nНовый документ сохранён:\n{dst}")
        except Exception as exc: messagebox.showerror(APP_TITLE, f"Не удалось оформить документ:\n{exc}")

def main():
    root = Tk(); App(root); root.mainloop()

if __name__ == "__main__": main()
