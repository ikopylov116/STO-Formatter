from pathlib import Path
from docx import Document

from app.sto.p305_2026 import P3052026Profile
from app.document.checker import check_document
from app.document.formatting import format_document


def make_doc(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Введение")
    doc.add_paragraph("Это обычный текст документа.")
    doc.save(path)


def test_format_creates_new_file(tmp_path):
    source = tmp_path / "source.docx"
    target = tmp_path / "formatted.docx"
    make_doc(source)
    profile = P3052026Profile()
    findings = format_document(str(source), str(target), profile.config, profile.structural_elements)
    assert target.exists()
    assert any(f.fixed for f in findings)
    assert source.read_bytes() != target.read_bytes()


def test_formatted_document_has_expected_margins_and_page_field(tmp_path):
    source = tmp_path / "source.docx"
    target = tmp_path / "formatted.docx"
    make_doc(source)
    profile = P3052026Profile()
    format_document(str(source), str(target), profile.config, profile.structural_elements)
    doc = Document(target)
    section = doc.sections[0]
    assert round(section.left_margin.cm, 2) == 3.0
    assert round(section.right_margin.cm, 2) == 1.0
    assert round(section.top_margin.cm, 2) == 2.0
    assert round(section.bottom_margin.cm, 2) == 2.0
    assert " PAGE " in section.footer._element.xml


def test_checker_reports_double_spaces(tmp_path):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Текст  с двойным пробелом")
    doc.save(path)
    profile = P3052026Profile()
    findings = check_document(str(path), profile.config, profile.structural_elements)
    finding = next(f for f in findings if f.rule_id == "P305-4.1.13")
    assert finding.severity.value == "warning"


def test_profile_id():
    assert P3052026Profile().id == "P_305.2026"
